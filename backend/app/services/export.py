# -*- coding: utf-8 -*-
"""P3-3 导出三件套：意见书 PDF + Excel 明细 + 高管摘要（批注 docx 已有）。

- opinion_pdf：reportlab 生成《合同审查意见书》，中文用内置 CID 字体 STSong-Light；
- findings_xlsx：openpyxl 生成问题明细 Excel（严重度/检查器/规则/说明/原文定位/建议）；
- executive_summary：规则化生成高管摘要（无 LLM、0 token）：总体结论 + 严重度分布 +
  Top 风险 + 放行建议，随 PDF 首部输出。
"""

from __future__ import annotations

import io
from collections import Counter
from datetime import datetime
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_SEVERITY_ORDER = {"error": 0, "warning": 1, "info": 2}
_SEVERITY_LABEL = {"error": "错误", "warning": "警告", "info": "提示"}


def _severity_sort_key(f) -> int:
    return _SEVERITY_ORDER.get(f.severity, 3)


def executive_summary(doc, findings: list) -> dict:
    """规则化高管摘要：数字全部来自 findings，无模型调用。"""
    total = len(findings)
    by_sev = Counter(f.severity for f in findings)
    errors = by_sev.get("error", 0)
    warnings = by_sev.get("warning", 0)
    infos = by_sev.get("info", 0)

    # Top 风险：按 rule_id 聚类取前 3
    by_rule: Counter = Counter(f.rule_id for f in findings)
    top_rules = by_rule.most_common(3)

    if errors == 0 and warnings == 0:
        conclusion = "未发现错误与警告类问题，文档整体质量良好，可进入正常流程。"
        advice = "建议放行"
    elif errors == 0:
        conclusion = f"发现 {warnings} 处警告，建议核实后处理，无阻断性错误。"
        advice = "核实警告后放行"
    else:
        conclusion = f"发现 {errors} 处错误级问题，存在实质性风险，需修改后重新校验。"
        advice = "暂不放行，需修改复验"

    return {
        "doc_filename": doc.filename if doc else "-",
        "total": total,
        "errors": errors,
        "warnings": warnings,
        "infos": infos,
        "top_rules": [{"rule_id": r, "count": c} for r, c in top_rules],
        "conclusion": conclusion,
        "advice": advice,
    }


def _init_cid_font() -> None:
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    except Exception:
        pass


def _pdf_escape(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def opinion_pdf_bytes(doc, findings: list) -> bytes:
    """《合同审查意见书》PDF（A4，中文 CID 字体，含高管摘要）。"""
    _init_cid_font()
    summary = executive_summary(doc, findings)
    buf = io.BytesIO()
    doc_pdf = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
        title="合同审查意见书",
    )
    title_style = ParagraphStyle("title", fontName="STSong-Light", fontSize=18,
                                 leading=26, alignment=1, spaceAfter=6 * mm)
    h_style = ParagraphStyle("h", fontName="STSong-Light", fontSize=12,
                             leading=18, spaceBefore=4 * mm, spaceAfter=2 * mm)
    body = ParagraphStyle("body", fontName="STSong-Light", fontSize=10.5,
                          leading=16, wordWrap="CJK")

    story = [
        Paragraph("合同审查意见书", title_style),
        Paragraph(f"审查时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}"
                  f"　文档：{_pdf_escape(summary['doc_filename'])}", body),
        Spacer(1, 4 * mm),
        Paragraph("一、总体结论（高管摘要）", h_style),
        Paragraph(_pdf_escape(summary["conclusion"]), body),
        Paragraph(f"发现问题 {summary['total']} 处：错误 {summary['errors']}、"
                  f"警告 {summary['warnings']}、提示 {summary['infos']}。", body),
        Paragraph("处理建议：" + _pdf_escape(summary["advice"]), body),
        Spacer(1, 3 * mm),
        Paragraph("二、风险分布", h_style),
    ]
    if summary["top_rules"]:
        for r in summary["top_rules"]:
            story.append(Paragraph(f"· {_pdf_escape(r['rule_id'])}：{r['count']} 处", body))
    else:
        story.append(Paragraph("无高发规则。", body))

    story.append(Spacer(1, 3 * mm))
    story.append(Paragraph("三、问题清单", h_style))
    rows = [["级别", "规则", "说明", "建议"]]
    for f in sorted(findings, key=_severity_sort_key):
        rows.append([
            _SEVERITY_LABEL.get(f.severity, f.severity),
            _pdf_escape(f.rule_id),
            _pdf_escape((f.message or "")[:120]),
            _pdf_escape((f.suggestion or "")[:80]),
        ])
    if len(rows) > 1:
        table = Table(rows, colWidths=[16 * mm, 40 * mm, 78 * mm, 46 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
            ("FONTSIZE", (0, 0), (-1, -1), 8.5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(table)
    else:
        story.append(Paragraph("未发现问题。", body))

    doc_pdf.build(story)
    return buf.getvalue()


def findings_xlsx_bytes(doc, findings: list) -> bytes:
    """问题明细 Excel（openpyxl）：工作表含文档信息 + 逐条明细。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "问题明细"
    header = ["级别", "检查器", "规则", "说明", "建议", "状态", "原文定位"]
    ws.append(header)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4472C4")
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sev_fill = {"error": PatternFill("solid", fgColor="FFC7CE"),
                "warning": PatternFill("solid", fgColor="FFEB9C"),
                "info": PatternFill("solid", fgColor="C6EFCE")}
    for f in sorted(findings, key=_severity_sort_key):
        ws.append([
            _SEVERITY_LABEL.get(f.severity, f.severity),
            getattr(f, "checker", "") or "",
            f.rule_id,
            f.message,
            f.suggestion,
            getattr(f, "status", "") or "",
            getattr(f, "anchor", "") or "",
        ])
        row = ws.max_row
        fill = sev_fill.get(f.severity)
        if fill:
            ws.cell(row=row, column=1).fill = fill

    for col, width in zip(range(1, len(header) + 1), [10, 12, 22, 60, 45, 12, 30]):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("汇总")
    summary = executive_summary(doc, findings)
    ws2.append(["合同审查意见书汇总"])
    ws2["A1"].font = Font(bold=True, size=14)
    for k, v in [
        ("文档", summary["doc_filename"]),
        ("审查时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("问题总数", summary["total"]),
        ("错误", summary["errors"]),
        ("警告", summary["warnings"]),
        ("提示", summary["infos"]),
        ("总体结论", summary["conclusion"]),
        ("处理建议", summary["advice"]),
    ]:
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 70

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
