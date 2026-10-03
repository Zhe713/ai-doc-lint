"""annotate 批注版任务测试：docx 原生批注、未匹配汇总、问题清单、无问题拦截、归属隔离。"""

import io
import uuid
import zipfile
from types import SimpleNamespace

from docx import Document as DocxDocument

from app.db import storage
from app.models.document import (
    ART_ANNOTATED,
    Artifact,
    CheckFinding,
    Document,
    DocumentTask,
)
from conftest import TestSessionLocal
from sqlalchemy.orm import Session
from tests.test_intake import _auth, _make_user, _token


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _doc_text_from_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _setup(*, ext="docx", original: bytes | None, findings: list[dict],
           username=None):
    db = TestSessionLocal()
    username = username or f"ann_{uuid.uuid4().hex[:8]}"
    _make_user(username)
    from app.models.user import User

    user = db.query(User).filter(User.username == username).one()

    orig_key = f"docs/{user.id}/orig.{ext}"
    storage.put_object(orig_key, original or b"",
                       content_type="application/octet-stream")
    doc = Document(
        user_id=user.id, filename=f"sample.{ext}", storage_key=orig_key,
        sha256="x" * 64, size_bytes=len(original or b""),
        content_type="text/plain", ext=ext, doc_type="contract",
        status="parsed",
    )
    db.add(doc)
    db.flush()

    for spec in findings:
        db.add(CheckFinding(
            document_id=doc.id, user_id=user.id,
            checker=spec.get("checker", "clause"),
            rule_id=spec.get("rule_id", "clause.breach"),
            severity=spec.get("severity", "warning"),
            location=spec.get("location"),
            message=spec["message"],
            suggestion=spec.get("suggestion"),
            status=spec.get("status", "open"),
        ))

    task = DocumentTask(
        group_id=None, document_id=doc.id, user_id=user.id,
        task_type="annotate", status="pending", max_attempts=1,
    )
    db.add(task)
    db.commit()
    return db, doc, task


def _run_task(task: DocumentTask):
    from app.tasks import annotate as ann_mod

    # 让任务函数使用测试库会话（默认 SessionLocal 指向 dev 库）
    ann_mod.SessionLocal = TestSessionLocal
    ann_mod.run_annotate.run(str(task.id))


def _output_bytes(db: Session, task: DocumentTask) -> bytes:
    db.refresh(task)
    key = task.result_refs["annotated"]
    return storage.get_object_bytes(key)


def test_docx_native_comment():
    para = "甲方应在约定时间内交付合格货物，乙方按时验收。"
    original = _docx_bytes([para])
    db, doc, task = _setup(
        ext="docx", original=original,
        findings=[{
            "rule_id": "clause.performance", "severity": "warning",
            "message": "履行期限约定不明确",
            "suggestion": "明确具体交付日期",
            "location": {"snippet": "甲方应在约定时间内交付合格货物"},
        }],
    )
    _run_task(task)
    data = _output_bytes(db, task)

    zf = zipfile.ZipFile(io.BytesIO(data))
    assert "word/comments.xml" in zf.namelist()
    comments = zf.read("word/comments.xml").decode("utf-8")
    assert "履行期限约定不明确" in comments
    assert "明确具体交付日期" in comments
    assert db.query(Artifact).filter(
        Artifact.document_id == doc.id, Artifact.kind == ART_ANNOTATED
    ).count() == 1
    db.close()


def test_unmatched_finding_goes_summary():
    original = _docx_bytes(["这是一段没有任何定位线索的正文内容。"])
    db, doc, task = _setup(
        ext="docx", original=original,
        findings=[{
            "message": "缺少违约责任条款",
            "location": {"snippet": "原文里完全没有的句子"},
        }],
    )
    _run_task(task)
    text = _doc_text_from_docx(_output_bytes(db, task))
    assert "问题汇总" in text
    assert "缺少违约责任条款" in text
    db.close()


def test_checklist_for_pdf_txt():
    db, doc, task = _setup(
        ext="txt", original=b"plain text body",
        findings=[
            {"message": "缺少争议解决条款", "severity": "warning"},
            {"message": "缺少违约责任条款", "severity": "warning"},
        ],
    )
    _run_task(task)
    text = _doc_text_from_docx(_output_bytes(db, task))
    assert "问题清单" in text
    assert "缺少争议解决条款" in text
    assert "缺少违约责任条款" in text
    db.close()


def test_no_open_findings_fails():
    original = _docx_bytes(["正文内容。"])
    db, doc, task = _setup(
        ext="docx", original=original,
        findings=[{"message": "已处理的问题", "status": "fixed"}],
    )
    _run_task(task)
    db.refresh(task)
    assert task.status == "failed"
    assert "无需生成" in (task.error_message or "")
    db.close()


def test_annotate_api_ownership(client, monkeypatch):
    monkeypatch.setattr("app.api.documents.enqueue_parse", lambda task_id: None)
    headers = _auth(client, "ann_owner")
    resp = client.post(
        "/api/v1/documents/upload",
        files={"files": ("a.txt", b"hello", "text/plain")},
        headers=headers,
    )
    doc_id = resp.json()["tasks"][0]["document_id"]

    created = client.post(f"/api/v1/documents/{doc_id}/annotate", headers=headers)
    assert created.status_code == 201

    _make_user("ann_other")
    other = {"Authorization": f"Bearer {_token(client, 'ann_other')}"}
    resp = client.post(f"/api/v1/documents/{doc_id}/annotate", headers=other)
    assert resp.status_code == 404
