# -*- coding: utf-8 -*-
"""P3-3 导出三件套测试：高管摘要规则化 + 意见书 PDF + 明细 Excel。"""

import io
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdfium2 import PdfDocument

from app.main import app
from app.services.export import executive_summary, findings_xlsx_bytes, opinion_pdf_bytes

client = TestClient(app)


class _F:
    def __init__(self, severity, rule_id, message, suggestion, checker="format", status="open", anchor=""):
        self.severity = severity
        self.rule_id = rule_id
        self.message = message
        self.suggestion = suggestion
        self.checker = checker
        self.status = status
        self.anchor = anchor


class _D:
    def __init__(self, filename="测试合同.pdf"):
        self.filename = filename


def test_executive_summary_clean():
    s = executive_summary(_D(), [])
    assert s["total"] == 0
    assert "未发现" in s["conclusion"]
    assert s["advice"] == "建议放行"


def test_executive_summary_error_block():
    s = executive_summary(_D(), [_F("error", "format.x", "金额大小写不一致", "建议修改")])
    assert s["errors"] == 1
    assert s["advice"] == "暂不放行，需修改复验"


def test_executive_summary_warning_only():
    s = executive_summary(_D(), [_F("warning", "typo.x", "疑似错别字", "")])
    assert s["warnings"] == 1 and s["errors"] == 0
    assert s["advice"] == "核实警告后放行"


def test_opinion_pdf_contains_summary():
    findings = [_F("error", "format.amount", "金额大小写不一致", "建议修改为人民币壹万元整"),
                _F("warning", "typo.word", "疑似错别字：由其", "建议改为尤其")]
    raw = opinion_pdf_bytes(_D(), findings)
    assert raw.startswith(b"%PDF")
    pdf = PdfDocument(raw)
    assert len(pdf) >= 1
    text = ""
    for page in pdf:
        text += page.get_textpage().get_text_bounded()
    assert "合同审查意见书" in text
    assert "暂不放行" in text
    assert "2 处" in text or "错误 1" in text


def test_opinion_pdf_empty():
    raw = opinion_pdf_bytes(_D(), [])
    assert raw.startswith(b"%PDF")
    pdf = PdfDocument(raw)
    text = "".join(p.get_textpage().get_text_bounded() for p in pdf)
    assert "未发现问题" in text


def test_findings_xlsx_shape():
    findings = [_F("error", "format.amount", "金额不一致", "修改", anchor="1.2"),
                _F("info", "field.contact", "联系电话为空", "", anchor="3.1")]
    raw = findings_xlsx_bytes(_D("钢材买卖合同.pdf"), findings)
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["问题明细"]
    assert ws.max_row == 3  # 表头 + 2 条
    assert ws.cell(row=2, column=1).value == "错误"
    ws2 = wb["汇总"]
    assert ws2["A2"].value == "文档"
    assert ws2["B2"].value == "钢材买卖合同.pdf"


def test_export_pdf_requires_login():
    resp = client.get("/api/v1/documents/00000000-0000-0000-0000-000000000000/export/pdf")
    assert resp.status_code in (401, 403)


def test_export_xlsx_requires_login():
    resp = client.get("/api/v1/documents/00000000-0000-0000-0000-000000000000/export/xlsx")
    assert resp.status_code in (401, 403)
