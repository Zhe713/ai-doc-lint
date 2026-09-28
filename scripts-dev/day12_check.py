"""Day12 验证：构造一好一坏两份 docx，验证格式检查器的命中与不误报。"""
import io
import time

import httpx
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

BASE = "http://localhost:8000/api/v1"


def set_cjk(run, font):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)


def add_body(doc, text, *, font="宋体", size=12, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_cjk(run, font)
    run.font.size = Pt(size)
    return p


def build_good():
    doc = DocxDocument()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.17)
    doc.add_paragraph("一、标的")
    doc.add_paragraph("（一）货物名称")
    doc.add_paragraph("1. 品名规格")
    add_body(doc, "本合同项下货物为标准办公设备一批，具体型号与数量以附件清单为准。")
    add_body(doc, "货物应当符合国家质量标准并随附合格证明，包装适用于长途运输。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_bad():
    doc = DocxDocument()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(216), Mm(279)  # Letter
    sec.top_margin = Cm(6)  # 超区间
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(3.0)
    doc.add_paragraph("一、标的")
    doc.add_paragraph("1. 品名规格")  # 跳级，缺（一）
    add_body(doc, "本合同项下货物为标准办公设备一批，具体型号与数量以双方确认为准。",
             font="微软雅黑", size=20, indent=False)
    add_body(doc, "货物应当符合国家质量标准并随附合格证明，包装应当适用于长途运输。",
             font="微软雅黑", size=20, indent=False)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def upload_and_wait(name, content):
    r = httpx.post(BASE + "/auth/login",
                   json={"username": "boss_admin", "password": "AdminPass123"})
    headers = {"Authorization": f"Bearer {r.json()['access_token']}"}
    r = httpx.post(
        BASE + "/documents/upload", headers=headers,
        files={"files": (name, io.BytesIO(content),
                         "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    doc_id = r.json()["tasks"][0]["document_id"]

    from app.db.session import SessionLocal
    from app.models.document import CheckFinding, DocumentTask

    for _ in range(40):
        db = SessionLocal()
        review = (
            db.query(DocumentTask)
            .filter(DocumentTask.document_id == doc_id, DocumentTask.task_type == "review")
            .one_or_none()
        )
        if review and review.status in ("parsed", "failed"):
            findings = db.query(CheckFinding).filter(CheckFinding.document_id == doc_id).all()
            data = [(f.rule_id, f.severity, f.message) for f in findings]
            stats = review.stats
            db.close()
            return data, stats
        db.close()
        time.sleep(1)
    raise AssertionError("review timeout")


good, gstats = upload_and_wait("day12_good.docx", build_good())
print("GOOD findings:", good)
print("GOOD stats:", gstats)

bad, bstats = upload_and_wait("day12_bad.docx", build_bad())
print("BAD findings:")
for rule, sev, msg in bad:
    print(f"  [{sev}] {rule}: {msg}")
print("BAD stats:", bstats)

assert len(good) == 0, f"好文档不应有问题，实际: {good}"
bad_rules = {rule for rule, _, _ in bad}
for expected in ("format.page.size", "format.page.margin", "format.font.body",
                 "format.font.size", "format.paragraph.indent", "format.heading.sequence"):
    assert expected in bad_rules, f"坏文档缺少命中 {expected}"
print("DAY12 OK")
